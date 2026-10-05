"""Validated single-chain sequence substitution and peptide-window extraction.

No structure modeling, energy calculation, or epitope prediction is performed.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import platform
import re
import sys
from dataclasses import dataclass
from pathlib import Path

__version__ = '1.0.0'
AA = set('ACDEFGHIKLMNPQRSTVWY')
TOKEN = re.compile(r'^([A-Z])([A-Za-z0-9])(\d+)([A-Z])$')


class InputError(ValueError):
    """An input failed validation; no scientific results should be inferred."""


def sequence_text(text: str, label: str) -> str:
    seq = ''.join(text.split()).upper()
    if not seq or set(seq) - AA:
        raise InputError(f'{label}: expected a nonempty sequence of the 20 standard amino acids')
    if len(seq) > 32767:
        raise InputError(f'{label}: exceeds the Excel cell limit of 32767 characters')
    return seq


def read_sequence(path: Path) -> str:
    lines = path.read_text(encoding='utf-8-sig').splitlines()
    records = 0
    parts = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('>'):
            records += 1
            if records > 1 or parts:
                raise InputError('Supply exactly one FASTA record or one plain-text sequence')
        else:
            parts.append(line)
    return sequence_text(''.join(parts), 'Reference')


@dataclass(frozen=True)
class Mutation:
    token: str
    position: int  # One-based sequence index, after subtracting the offset.
    wt: str
    mutant: str


def parse_mutations(text: str, reference: str, chain: str = 'A', offset: int = 0) -> list[Mutation]:
    if not re.fullmatch(r'[A-Za-z0-9]', chain):
        raise InputError('Chain must be a single case-sensitive letter or digit')
    text = text.strip()
    if text.endswith(';'):
        text = text[:-1].strip()
    if text == 'WT':
        return []
    result = []
    seen = set()
    for raw in text.split(','):
        raw = raw.strip()
        match = TOKEN.fullmatch(raw)
        if not match:
            raise InputError(f'Invalid mutation token {raw!r}; expected AA1G (WT, chain, number, mutant)')
        wt, observed_chain, number, mutant = match.groups()
        if observed_chain != chain:
            raise InputError(f'{raw}: chain {observed_chain!r} differs from selected chain {chain!r}')
        if wt not in AA or mutant not in AA:
            raise InputError(f'{raw}: only standard amino acids are supported')
        pos = int(number) - offset
        if not 1 <= pos <= len(reference):
            raise InputError(f'{raw}: mapped sequence position {pos} is out of range')
        if pos in seen:
            raise InputError(f'{raw}: duplicate mutation position {pos}')
        if reference[pos - 1] != wt:
            raise InputError(f'{raw}: reference at sequence position {pos} is {reference[pos - 1]}, not {wt}')
        if wt == mutant:
            raise InputError(f'{raw}: a substitution must change the residue; use WT for an unchanged variant')
        seen.add(pos)
        result.append(Mutation(f'{wt}{chain}{int(number)}{mutant}', pos, wt, mutant))
    return sorted(result, key=lambda m: m.position)


def apply_mutations(reference: str, mutations: list[Mutation]) -> str:
    chars = list(reference)
    for m in mutations:
        chars[m.position - 1] = m.mutant
    return ''.join(chars)


def read_variants(path: Path, reference: str, chain: str, offset: int, max_mutations: int | None = None) -> list[dict]:
    variants = []
    for lineno, line in enumerate(path.read_text(encoding='utf-8-sig').splitlines(), 1):
        text = line.split('#', 1)[0].strip()
        if not text:
            continue
        try:
            mutations = parse_mutations(text, reference, chain, offset)
            if max_mutations is not None and len(mutations) > max_mutations:
                raise InputError(f'more than {max_mutations} mutations')
        except InputError as exc:
            raise InputError(f'Variants line {lineno}: {exc}') from exc
        variants.append({'variant_id': f'VAR{len(variants)+1:03d}',
                         'mutation_string': ','.join(m.token for m in mutations) or 'WT',
                         'n_mutations': len(mutations), 'validation': 'OK',
                         'mutant_sequence': apply_mutations(reference, mutations),
                         'mutations': mutations, 'source_line': lineno})
    if not variants:
        raise InputError('No variants found')
    return variants


def find_start(peptide: str, reference: str) -> int:
    first = reference.find(peptide)
    if first < 0:
        raise InputError('Native window does not occur in the reference')
    if reference.find(peptide, first + 1) >= 0:
        raise InputError('Native window occurs more than once; supply its explicit start coordinate')
    return first + 1


def read_windows(path: Path, reference: str) -> list[dict]:
    windows, keys, core_annotations = [], set(), {}
    with path.open(encoding='utf-8-sig', newline='') as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise InputError('Window CSV has no header')
        reader.fieldnames = [v.strip() for v in reader.fieldnames]
        required = {'core_id', 'window', 'window_len', 'native_window'}
        if len(set(reader.fieldnames)) != len(reader.fieldnames) or not required <= set(reader.fieldnames):
            raise InputError('Window CSV requires unique headers: core_id,window,window_len,native_window; optional start,core_seq')
        for lineno, row in enumerate(reader, 2):
            try:
                if None in row or any(value is None for value in row.values()):
                    raise InputError('Wrong number of columns')
                cid = row['core_id'].strip()
                if not cid or cid == 'NONE':
                    raise InputError('core_id must be nonempty and cannot be NONE')
                window, length = int(row['window']), int(row['window_len'])
                if window < 1 or length < 1:
                    raise InputError('Window number and length must be positive')
                native = sequence_text(row['native_window'], 'Native window')
                if len(native) != length:
                    raise InputError('window_len does not match native_window length')
                start_text = row.get('start', '').strip()
                start = int(start_text) if start_text else find_start(native, reference)
                end = start + length - 1
                if start < 1 or end > len(reference) or reference[start-1:end] != native:
                    raise InputError('Explicit window coordinates do not match the reference')
                key = (cid, window)
                if key in keys:
                    raise InputError('Duplicate core_id and window number')
                annotation = row.get('core_seq', '').strip()
                if annotation:
                    annotation = sequence_text(annotation, 'core_seq annotation')
                    if cid in core_annotations and core_annotations[cid] != annotation:
                        raise InputError('Inconsistent core_seq annotation within a core')
                    core_annotations[cid] = annotation
                keys.add(key)
                windows.append({'core_id': cid, 'window': window, 'window_len': length,
                                'native_window': native, 'start': start, 'end': end,
                                'core_seq': annotation})
            except (ValueError, KeyError) as exc:
                raise InputError(f'Windows line {lineno}: {exc}') from exc
    if not windows:
        raise InputError('No windows found')
    return sorted(windows, key=lambda w: (w['core_id'], w['window']))


VARIANT_FIELDS = ['variant_id', 'mutation_string', 'n_mutations', 'validation', 'mutant_sequence']
LEGACY_HEADERS = ['Variant ID', 'Mutation string', '# Mutations', 'Validation', 'Mutant sequence (mutated AA in red)']


def read_built_variants(path: Path, reference: str, chain: str, offset: int) -> list[dict]:
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            reader = csv.DictReader(handle)
            if not set(VARIANT_FIELDS) <= set(reader.fieldnames or []):
                raise InputError('Mutants CSV does not have the expected builder columns')
            records = list(reader)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=False)
        try:
            ws = wb['Mutant sequences'] if 'Mutant sequences' in wb.sheetnames else wb.active
            if [ws.cell(4, i).value for i in range(1, 6)] != LEGACY_HEADERS:
                raise InputError('Mutants workbook header must match the builder layout at row 4')
            records = [dict(zip(VARIANT_FIELDS, row[:5])) for row in ws.iter_rows(min_row=5, values_only=True)
                       if any(x is not None for x in row)]
        finally:
            wb.close()
    else:
        raise InputError('Mutants input must be a builder CSV or XLSX file')
    variants, ids = [], set()
    for index, record in enumerate(records, 1):
        try:
            vid = str(record.get('variant_id') or '').strip()
            if not vid or vid in ids:
                raise InputError('Missing or duplicate variant ID')
            if str(record.get('validation')).strip().upper() != 'OK':
                raise InputError('Variant is not marked OK; invalid rows cannot silently be skipped')
            mutations = parse_mutations(str(record.get('mutation_string') or ''), reference, chain, offset)
            if str(record.get('n_mutations')) != str(len(mutations)):
                raise InputError('Mutation count disagrees with the mutation string')
            seq = sequence_text(str(record.get('mutant_sequence') or ''), 'Mutant sequence')
            if seq != apply_mutations(reference, mutations):
                raise InputError('Mutant sequence does not match the supplied reference and substitutions')
            ids.add(vid)
            variants.append({'variant_id': vid, 'mutation_string': ','.join(m.token for m in mutations) or 'WT',
                             'n_mutations': len(mutations), 'validation': 'OK',
                             'mutant_sequence': seq, 'mutations': mutations})
        except (ValueError, TypeError) as exc:
            raise InputError(f'Mutants record {index}: {exc}') from exc
    if not variants:
        raise InputError('No built variants found')
    return variants


POOL_FIELDS = ['variant_id', 'mutation_string', 'core_id', 'window', 'start', 'end',
               'native_peptide', 'mutant_peptide', 'aa_changes', 'mutation_profile', 'status', 'peptide_id']


def generate_pools(variants: list[dict], windows: list[dict], selection: str = 'affected-core') -> tuple[list[dict], list[dict], list[dict]]:
    if selection not in {'affected-core', 'changed-windows'}:
        raise InputError('Unknown pool selection mode')
    rows, summaries = [], []
    for variant in variants:
        positions = {m.position for m in variant['mutations']}
        affected = {w['core_id'] for w in windows if any(w['start'] <= p <= w['end'] for p in positions)}
        selected = [w for w in windows if w['core_id'] in affected and
                    (selection == 'affected-core' or any(w['start'] <= p <= w['end'] for p in positions))]
        start_index = len(rows)
        for w in selected:
            peptide = variant['mutant_sequence'][w['start']-1:w['end']]
            changes = sum(a != b for a, b in zip(w['native_window'], peptide))
            profile = '+'.join(m.token for m in variant['mutations'] if w['start'] <= m.position <= w['end'])
            rows.append({'variant_id': variant['variant_id'], 'mutation_string': variant['mutation_string'],
                         'core_id': w['core_id'], 'window': w['window'], 'start': w['start'], 'end': w['end'],
                         'native_peptide': w['native_window'], 'mutant_peptide': peptide,
                         'aa_changes': changes, 'mutation_profile': profile or 'none',
                         'status': 'CHANGED' if changes else 'UNCHANGED', 'peptide_id': ''})
        changed = [r for r in rows[start_index:] if r['status'] == 'CHANGED']
        summaries.append({'variant_id': variant['variant_id'], 'mutation_string': variant['mutation_string'],
                          'affected_cores': ';'.join(sorted(affected)), 'window_rows': len(selected),
                          'changed_window_rows': len(changed), 'unique_changed_sequences': len({r['mutant_peptide'] for r in changed}),
                          'status': 'OK' if affected else 'NO_AFFECTED_WINDOWS'})
    unique_sequences = sorted({r['mutant_peptide'] for r in rows if r['status'] == 'CHANGED'})
    ids = {seq: f'PEP{i+1:04d}' for i, seq in enumerate(unique_sequences)}
    for row in rows:
        if row['status'] == 'CHANGED':
            row['peptide_id'] = ids[row['mutant_peptide']]
    catalog = []
    for seq in unique_sequences:
        members = [r for r in rows if r['status'] == 'CHANGED' and r['mutant_peptide'] == seq]
        catalog.append({'peptide_id': ids[seq], 'sequence': seq, 'length': len(seq),
                        'n_variants': len({r['variant_id'] for r in members}),
                        'variant_ids': ';'.join(sorted({r['variant_id'] for r in members})),
                        'window_sources': ';'.join(sorted({f"{r['core_id']}:{r['window']}" for r in members}))})
    return rows, summaries, catalog


def write_csv(path: Path, fields: list[str], rows: list[dict]) -> None:
    with path.open('w', encoding='utf-8', newline='') as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, extrasaction='ignore')
        writer.writeheader()
        writer.writerows(rows)


def rich_sequence(seq: str, positions: set[int]):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    groups = []
    for i, aa in enumerate(seq, 1):
        color = 'FFFF0000' if i in positions else 'FF000000'
        if groups and groups[-1][0] == color:
            groups[-1][1] += aa
        else:
            groups.append([color, aa])
    return CellRichText(*(TextBlock(InlineFont(color=color, sz=10), text) for color, text in groups))


def table_sheet(wb, title: str, fields: list[str], rows: list[dict], header_row: int = 1):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    ws = wb.create_sheet(title)
    for col, name in enumerate(fields, 1):
        cell = ws.cell(header_row, col, name)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='44546A')
        ws.column_dimensions[get_column_letter(col)].width = min(55, max(18, len(name) + 3))
    for rownum, row in enumerate(rows, header_row + 1):
        for col, field in enumerate(fields, 1):
            value = row.get(field, '')
            if isinstance(value, str) and (len(value) > 32767 or re.search(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', value)):
                raise InputError('Text exceeds Excel limits or contains unsupported control characters')
            cell = ws.cell(rownum, col, value)
            if isinstance(value, str):
                cell.data_type = 's'  # User labels must never become Excel formulas.
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = ws.cell(header_row + 1, 1)
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(fields))}{max(header_row, ws.max_row)}'
    return ws


def manifest(command: str, files: list[Path], chain: str, offset: int, extra: dict | None = None) -> dict:
    import openpyxl
    return {'schema_version': 1, 'tool_version': __version__, 'operation': command,
            'python': platform.python_version(), 'openpyxl': openpyxl.__version__,
            'chain': chain, 'position_offset': offset,
            'coordinate_rule': 'sequence_position = mutation_number - position_offset; windows are 1-based sequence coordinates',
            'inputs': [{'filename': p.name, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()} for p in files],
            **(extra or {})}


def write_metadata(wb, metadata: dict):
    table_sheet(wb, 'Run metadata', ['key', 'value'],
                [{'key': key, 'value': json.dumps(value, ensure_ascii=True)} for key, value in metadata.items()])


def build_outputs(args) -> None:
    from openpyxl import Workbook
    reference = read_sequence(args.sequence)
    variants = read_variants(args.variants, reference, args.chain, args.position_offset, args.max_mutations)
    metadata = manifest('mutants', [args.sequence, args.variants], args.chain, args.position_offset,
                        {'variant_count': len(variants), 'max_mutations': args.max_mutations})
    wb = Workbook()
    wb.remove(wb.active)
    xrows = [dict(zip(LEGACY_HEADERS, [v[k] for k in VARIANT_FIELDS])) for v in variants]
    ws = table_sheet(wb, 'Mutant sequences', LEGACY_HEADERS, xrows, 4)
    ws['A1'] = 'Validated mutant sequences'
    ws['A2'] = f'{len(variants)} variants; changed amino acids appear in red. Sequence construction only.'
    for i, v in enumerate(variants, 5):
        ws.cell(i, 5).value = rich_sequence(v['mutant_sequence'], {m.position for m in v['mutations']})
    write_metadata(wb, metadata)
    args.output.mkdir(parents=True, exist_ok=False)
    wb.save(args.output / 'mutant_sequences.xlsx')
    write_csv(args.output / 'mutant_sequences.csv', VARIANT_FIELDS, variants)
    with (args.output / 'mutant_sequences.fasta').open('w', encoding='utf-8') as handle:
        for v in variants:
            handle.write(f">{v['variant_id']} {v['mutation_string']}\n{v['mutant_sequence']}\n")
    (args.output / 'manifest.json').write_text(json.dumps(metadata, indent=2) + '\n', encoding='utf-8')
    print(f"Built {len(variants)} validated variants in {args.output}")


def pool_outputs(args) -> None:
    from openpyxl import Workbook
    reference = read_sequence(args.sequence)
    windows = read_windows(args.windows, reference)
    variants = read_built_variants(args.mutants, reference, args.chain, args.position_offset)
    rows, summaries, catalog = generate_pools(variants, windows, args.selection)
    metadata = manifest('pools', [args.sequence, args.windows, args.mutants], args.chain, args.position_offset,
                        {'selection': args.selection, 'variant_count': len(variants), 'catalog_count': len(catalog),
                         'catalog_policy': 'Changed peptides only; deduplicated globally by amino-acid sequence'})
    catalog_fields = ['peptide_id', 'sequence', 'length', 'n_variants', 'variant_ids', 'window_sources']
    window_fields = ['core_id', 'window', 'start', 'end', 'window_len', 'native_window', 'core_seq']
    wb = Workbook()
    wb.remove(wb.active)
    ws = table_sheet(wb, 'Peptide pools', POOL_FIELDS, rows)
    for index, row in enumerate(rows, 2):
        changed = {i for i, (a, b) in enumerate(zip(row['native_peptide'], row['mutant_peptide']), 1) if a != b}
        for field in ['native_peptide', 'mutant_peptide']:
            ws.cell(index, POOL_FIELDS.index(field) + 1).value = rich_sequence(row[field], changed)
    table_sheet(wb, 'Pool summary', list(summaries[0]), summaries)
    table_sheet(wb, 'Native windows', window_fields, windows)
    table_sheet(wb, 'Synthesis catalog', catalog_fields, catalog)
    write_metadata(wb, metadata)
    args.output.mkdir(parents=True, exist_ok=False)
    wb.save(args.output / 'peptide_pools.xlsx')
    for filename, fields, data in [('peptide_pools.csv', POOL_FIELDS, rows),
                                   ('pool_summary.csv', list(summaries[0]), summaries),
                                   ('native_windows.csv', window_fields, windows),
                                   ('synthesis_catalog.csv', catalog_fields, catalog)]:
        write_csv(args.output / filename, fields, data)
    (args.output / 'manifest.json').write_text(json.dumps(metadata, indent=2) + '\n', encoding='utf-8')
    print(f"Processed {len(variants)} variants: {len(rows)} window rows, {len(catalog)} unique changed peptides")
    n_empty = sum(v['status'] == 'NO_AFFECTED_WINDOWS' for v in summaries)
    if n_empty:
        print(f'{n_empty} variants have no affected windows; see Pool summary')
    print(f'Output: {args.output}')


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--version', action='version', version=__version__)
    sub = parser.add_subparsers(dest='command', required=True)
    for name in ['mutants', 'pools']:
        cmd = sub.add_parser(name)
        cmd.add_argument('--sequence', type=Path, required=True, help='One reference FASTA record or plain sequence')
        cmd.add_argument('--chain', default='A', help='Exact chain identifier expected in mutation tokens (default A)')
        cmd.add_argument('--position-offset', type=int, default=0, help='Subtract from token residue numbers to get 1-based sequence positions')
        cmd.add_argument('--output', type=Path, required=True, help='New output directory; existing directories are refused')
        if name == 'mutants':
            cmd.add_argument('--variants', type=Path, required=True)
            cmd.add_argument('--max-mutations', type=int, default=None, help='Optional cap per variant; no cap by default')
        else:
            cmd.add_argument('--mutants', type=Path, required=True, help='Builder CSV or XLSX')
            cmd.add_argument('--windows', type=Path, required=True)
            cmd.add_argument('--selection', choices=['affected-core', 'changed-windows'], default='affected-core')
    args = parser.parse_args(argv)
    try:
        if args.output.exists():
            raise InputError('Output directory already exists; choose a new directory')
        if args.command == 'mutants' and args.max_mutations is not None and args.max_mutations < 1:
            raise InputError('--max-mutations must be positive')
        if args.command == 'mutants':
            build_outputs(args)
        else:
            pool_outputs(args)
        return 0
    except (InputError, OSError, ImportError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
